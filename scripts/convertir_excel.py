#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
convertir_excel.py
Script de conversión del Excel 'Plan de estudios CE y ATI (1).xlsx'
a los CSV de entrada para el programa en C del proyecto CEmestre.

Genera:
    - data/catalogo_CE.csv
    - data/catalogo_ATI.csv
    - reportes/reporte_limpieza.txt

Autor: Integrante 1 - Módulo I/O
"""

import openpyxl
import csv
import os
import re
from datetime import datetime

# ============================================================
# CONSTANTES
# ============================================================

ARCHIVO_EXCEL = "Plan de estudios CE y ATI (1).xlsx"
DIR_DATA = "data"
DIR_REPORTES = "reportes"

# Mapa de días en español / abreviaturas -> letra única
MAPA_DIAS = {
    # Abreviaturas CE
    "LUN": "L", "MAR": "K", "MIE": "M", "JUE": "J", "VIE": "V", "SAB": "S",
    # Nombres completos ATI (con y sin tilde)
    "LUNES": "L", "MARTES": "K", "MIERCOLES": "M", "MIÉRCOLES": "M",
    "JUEVES": "J", "VIERNES": "V", "SABADO": "S", "SÁBADO": "S",
}

# Correcciones específicas de códigos (typos del Excel)
CORRECCIONES_CODIGOS = {
    "F1 1101": "FI-1101",
    "F1-1101": "FI-1101",
    "TI 4501": "TI-4500",
    "TI-4501": "TI-4500",
    "TI4501": "TI-4500",
    "MA2404": "MA-2404",
    "MA2405": "MA-2404",
    "MA2406": "MA-2404",
    "MA2407": "MA-2404",
    "MA2408": "MA-2404",
    "MA2409": "MA-2404",
}

# Correcciones específicas de nombres de cursos (typos del Excel)
CORRECCIONES_NOMBRES = {
    "Habilidades de comunicación en ingenieía": "Habilidades de Comunicación en Ingeniería",
    "Habilidades de comunicación en ingeniería": "Habilidades de Comunicación en Ingeniería",
    "Comuncación oral": "Comunicación Oral",
    "Comunicación oral": "Comunicación Oral",
    "Exámen diagnóstico": "Examen Diagnóstico",
    "Examen diagnóstico": "Examen Diagnóstico",
    "Modelos organizacionales y gestion de TI": "Modelos Organizacionales y Gestión de TI",
    "Modelos organizacionales y gestión de TI": "Modelos Organizacionales y Gestión de TI",
    "Circuitos eléctricos en corriente contínua": "Circuitos Eléctricos en Corriente Continua",
    "Circuitos eléctricos en corriente continua": "Circuitos Eléctricos en Corriente Continua",
}

# Encabezado del CSV
CSV_HEADER = [
    "codigo", "nombre", "creditos", "semestre",
    "requisitos", "correquisitos",
    "grupo", "profesor", "dia", "hora_inicio", "hora_fin"
]

# ============================================================
# ESTADO GLOBAL
# ============================================================

reporte_lineas = []

def log(mensaje, nivel="INFO"):
    """Agrega una línea al reporte y la imprime en consola."""
    prefijo = {
        "INFO": "[INFO]",
        "WARN": "[WARN]",
        "ERROR": "[ERROR]",
        "OK": "[OK]",
    }.get(nivel, "[INFO]")
    linea = f"{prefijo} {mensaje}"
    reporte_lineas.append(linea)
    print(linea)

# ============================================================
# NORMALIZACIÓN
# ============================================================

def normalizar_codigo(codigo_raw):
    """
    Convierte 'CE 1101', 'CE1101', 'CE-1101' -> 'CE-1101'.
    Aplica correcciones específicas primero.
    """
    if codigo_raw is None:
        return None
    codigo = str(codigo_raw).strip()

    # Correcciones específicas (typos conocidos)
    if codigo in CORRECCIONES_CODIGOS:
        log(f"Corrección de código: '{codigo}' -> '{CORRECCIONES_CODIGOS[codigo]}'", "WARN")
        return CORRECCIONES_CODIGOS[codigo]

    # Patrón general: 2 letras + número de 4 dígitos (con o sin guion/espacio)
    match = re.match(r"^([A-Z]{2})\s*-?\s*(\d{4})$", codigo.upper())
    if match:
        return f"{match.group(1)}-{match.group(2)}"

    # Si no matchea, devolver tal cual (y reportar)
    log(f"Código no reconocido: '{codigo}'", "WARN")
    return codigo

def normalizar_dia(dia_raw):
    """Convierte 'MAR', 'Lunes', 'MIE' -> 'K', 'L', 'M'."""
    if dia_raw is None:
        return None
    dia = str(dia_raw).strip().upper()
    return MAPA_DIAS.get(dia)

def normalizar_hora(hora_raw):
    """
    Convierte '7:30' -> '07:30'.
    Retorna None si no se puede parsear.
    """
    if hora_raw is None:
        return None
    hora = str(hora_raw).strip()
    match = re.match(r"^(\d{1,2}):(\d{2})$", hora)
    if match:
        hh = int(match.group(1))
        mm = int(match.group(2))
        if 0 <= hh <= 23 and 0 <= mm <= 59:
            return f"{hh:02d}:{mm:02d}"
    log(f"Hora no reconocida: '{hora}'", "WARN")
    return None

def normalizar_profesor(nombre_raw):
    """
    Convierte 'SCHMIDT PERALTA JEFF' -> 'Schmidt Peralta Jeff'.
    Mantiene tildes.
    """
    if nombre_raw is None:
        return None
    nombre = str(nombre_raw).strip()
    if not nombre:
        return None
    # Capitalizar cada palabra manteniendo tildes
    return " ".join(p.capitalize() for p in nombre.split())

def normalizar_nombre_curso(nombre_raw):
    """
    Limpia el nombre del curso:
    - Quita el paréntesis con créditos al final: 'Introducción (3)' -> 'Introducción'
    - Aplica correcciones específicas
    """
    if nombre_raw is None:
        return None
    nombre = str(nombre_raw).strip()

    # Quitar el paréntesis con créditos al final
    nombre = re.sub(r"\s*\(\s*\d+\s*\)\s*$", "", nombre).strip()

    # Aplicar correcciones específicas
    if nombre in CORRECCIONES_NOMBRES:
        nombre = CORRECCIONES_NOMBRES[nombre]

    return nombre

def separar_lista_codigos(texto):
    """
    Separa 'CE 1101, CE 1104, MA 1403' -> ['CE-1101', 'CE-1104', 'MA-1403'].
    También separa por 2+ espacios consecutivos.
    """
    if texto is None:
        return []
    t = str(texto).strip()
    if not t or t == "-":
        return []
    # Separar por coma, salto de línea, punto y coma, o 2+ espacios
    partes = re.split(r"[,\n;]+|\s{2,}", t)
    resultado = []
    for p in partes:
        p = p.strip()
        if not p or p == "-":
            continue
        # Ignorar anotaciones tipo "(Correquisito)"
        p = re.sub(r"\(Correquisito\)", "", p, flags=re.IGNORECASE).strip()
        if p:
            codigo = normalizar_codigo(p)
            if codigo and codigo not in resultado:
                resultado.append(codigo)
    return resultado

def separar_requisitos_y_correquisitos(texto):
    """
    Dada una celda con mezcla, separa requisitos normales y correquisitos.
    Retorna (requisitos, correquisitos).
    Soporta separación por coma, salto de línea, punto y coma, o 2+ espacios.
    """
    if texto is None:
        return [], []
    t = str(texto).strip()
    if not t or t == "-":
        return [], []

    reqs = []
    correqs = []

    # Separar por coma, salto de línea, punto y coma, o 2+ espacios
    partes = re.split(r"[,\n;]+|\s{2,}", t)
    for p in partes:
        p = p.strip()
        if not p or p == "-":
            continue
        es_correq = "(Correquisito)" in p or "(correquisito)" in p.lower()
        codigo_limpio = re.sub(r"\(Correquisito\)", "", p, flags=re.IGNORECASE).strip()
        if not codigo_limpio:
            continue
        codigo = normalizar_codigo(codigo_limpio)
        if codigo:
            if es_correq:
                if codigo not in correqs:
                    correqs.append(codigo)
            else:
                if codigo not in reqs:
                    reqs.append(codigo)
    return reqs, correqs

# ============================================================
# LECTURA DEL PLAN DE ESTUDIOS
# ============================================================

def leer_plan_estudios(wb, nombre_hoja, fila_inicio):
    """
    Lee la hoja del plan de estudios (CE o ATI).
    Retorna un diccionario {codigo: {nombre, creditos, semestre, requisitos, correquisitos}}
    """
    ws = wb[nombre_hoja]
    cursos = {}

    # Columnas: B=Bloque0, C=Req0, D=Bloque1, E=Req1, ..., K=Req4
    for fila in ws.iter_rows(min_row=fila_inicio, values_only=True):
        for i in range(1, 11, 2):  # 1, 3, 5, 7, 9
            celda_curso = fila[i] if i < len(fila) else None
            celda_reqs = fila[i + 1] if i + 1 < len(fila) else None

            if celda_curso is None:
                continue

            curso_txt = str(celda_curso).strip()
            if not curso_txt or curso_txt == "-":
                continue

            # Parsear: "CE 1101 - Introducción a la programación (3)"
            match = re.match(r"^(.+?)\s*-\s*(.+?)\s*\(\s*(\d+)\s*\)\s*$", curso_txt)
            if not match:
                log(f"No se pudo parsear curso: '{curso_txt}'", "WARN")
                continue

            codigo_raw = match.group(1).strip()
            nombre_raw = match.group(2).strip()
            creditos = int(match.group(3))

            codigo = normalizar_codigo(codigo_raw)
            nombre = normalizar_nombre_curso(nombre_raw)

            # Semestre: según el índice del par
            semestre = (i - 1) // 2

            reqs, correqs = separar_requisitos_y_correquisitos(celda_reqs)

            cursos[codigo] = {
                "codigo": codigo,
                "nombre": nombre,
                "creditos": creditos,
                "semestre": semestre,
                "requisitos": reqs,
                "correquisitos": correqs,
            }

    log(f"Hoja '{nombre_hoja}': {len(cursos)} cursos leídos", "OK")
    return cursos

# ============================================================
# LECTURA DE HORARIOS
# ============================================================

def parsear_horario_CE(texto):
    """
    Parsea 'MAR[07:30-09:20] JUE[07:30-09:20]' -> [('K','07:30','09:20'), ('J','07:30','09:20')]
    """
    if texto is None:
        return []
    t = str(texto).strip()
    if not t:
        return []

    bloques = []
    matches = re.findall(r"([A-Z]+)\[(\d{1,2}:\d{2})-(\d{1,2}:\d{2})\]", t)
    for dia_raw, hi_raw, hf_raw in matches:
        dia = normalizar_dia(dia_raw)
        hi = normalizar_hora(hi_raw)
        hf = normalizar_hora(hf_raw)
        if dia and hi and hf:
            bloques.append((dia, hi, hf))
        else:
            log(f"Bloque horario no parseable: {dia_raw}[{hi_raw}-{hf_raw}]", "WARN")
    return bloques

def parsear_horario_ATI(texto):
    """
    Parsea 'Lunes - 18:00 , 20:50' -> [('L','18:00','20:50')]
    Soporta variantes con errores.
    NO emite warnings: si falla, retorna [] y el caller prueba otro formato.
    """
    if texto is None:
        return []
    t = str(texto).strip()
    if not t:
        return []

    # Normalizar separadores problemáticos
    # Caso 'Miercoles - 17:0:20:50' -> 'Miercoles - 17:00 , 20:50'
    t = re.sub(r"(\d{1,2}):(\d):(\d{2}):(\d{2})", r"\1:0\2 , \3:\4", t)
    # Caso 'Miercoles - 15:00 18:50' -> 'Miercoles - 15:00 , 18:50'
    t = re.sub(r"(\d{1,2}:\d{2})\s+(\d{1,2}:\d{2})", r"\1 , \2", t)

    match = re.match(
        r"^([A-Za-zÁÉÍÓÚáéíóúñÑ]+)\s*-\s*(\d{1,2}:\d{2})\s*,\s*(\d{1,2}:\d{2})$",
        t
    )
    if match:
        dia_raw, hi_raw, hf_raw = match.groups()
        dia = normalizar_dia(dia_raw)
        hi = normalizar_hora(hi_raw)
        hf = normalizar_hora(hf_raw)
        if dia and hi and hf:
            return [(dia, hi, hf)]
    return []

def leer_horarios(wb, nombre_hoja, fila_inicio, es_CE):
    """
    Lee la hoja de horarios.
    Retorna {codigo: {grupo: {'profesor': str, 'horarios': [(dia, hi, hf)]}}}
    """
    ws = wb[nombre_hoja]
    cursos = {}

    for fila in ws.iter_rows(min_row=fila_inicio, values_only=True):
        codigo_raw = fila[1] if len(fila) > 1 else None
        grupo_raw = fila[3] if len(fila) > 3 else None
        horario_raw = fila[4] if len(fila) > 4 else None
        profesor_raw = fila[6] if len(fila) > 6 else None

        if codigo_raw is None or grupo_raw is None:
            continue

        codigo = normalizar_codigo(codigo_raw)
        try:
            grupo = int(grupo_raw)
        except (ValueError, TypeError):
            log(f"Grupo no numérico: '{grupo_raw}' en {codigo}", "WARN")
            continue

        profesor = normalizar_profesor(profesor_raw)

        # Parsear horario según carrera
        if es_CE:
            bloques = parsear_horario_CE(horario_raw)
        else:
            # ATI: intentar formato ATI primero, si falla probar formato CE
            # (algunos cursos compartidos como MA-* vienen en formato CE)
            bloques = parsear_horario_ATI(horario_raw)
            if not bloques:
                bloques = parsear_horario_CE(horario_raw)

        if not bloques:
            log(f"Sin bloques horarios válidos: {codigo} grupo {grupo} -> {horario_raw!r}", "WARN")
            continue

        # Inicializar estructura
        if codigo not in cursos:
            cursos[codigo] = {}

        # Consolidar grupos (mismo grupo, mismo profesor -> agregar días)
        if grupo in cursos[codigo]:
            grupo_data = cursos[codigo][grupo]
            for b in bloques:
                if b not in grupo_data["horarios"]:
                    grupo_data["horarios"].append(b)
        else:
            cursos[codigo][grupo] = {
                "profesor": profesor,
                "horarios": list(bloques),
            }

    log(f"Hoja '{nombre_hoja}': {len(cursos)} cursos con horarios leídos", "OK")
    return cursos

# ============================================================
# ESCRITURA DEL CSV
# ============================================================

def escribir_csv(ruta, plan, horarios, cursos_sin_horario=None):
    """
    Escribe el CSV combinando plan de estudios + horarios.
    Una fila por bloque horario.
    """
    cursos_sin_horario = cursos_sin_horario or set()
    filas = []

    for codigo, info in plan.items():
        if codigo in horarios:
            grupos = horarios[codigo]
            for num_grupo in sorted(grupos.keys()):
                gdata = grupos[num_grupo]
                for dia, hi, hf in gdata["horarios"]:
                    filas.append([
                        info["codigo"],
                        info["nombre"],
                        info["creditos"],
                        info["semestre"],
                        ";".join(info["requisitos"]),
                        ";".join(info["correquisitos"]),
                        num_grupo,
                        gdata["profesor"] or "",
                        dia,
                        hi,
                        hf,
                    ])
        else:
            # Curso sin horario (diagnóstico)
            filas.append([
                info["codigo"],
                info["nombre"],
                info["creditos"],
                info["semestre"],
                ";".join(info["requisitos"]),
                ";".join(info["correquisitos"]),
                "", "", "", "", "",
            ])

    # Ordenar por código, luego grupo, luego día
    filas.sort(key=lambda f: (f[0], str(f[6]), str(f[8])))

    # Escribir
    with open(ruta, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(CSV_HEADER)
        writer.writerows(filas)

    log(f"CSV escrito: {ruta} ({len(filas)} filas)", "OK")
    return len(filas)

# ============================================================
# MAIN
# ============================================================

def main():
    log("=" * 60)
    log("CONVERSIÓN EXCEL -> CSV", "INFO")
    log(f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    log("=" * 60)

    # Crear carpetas de salida
    os.makedirs(DIR_DATA, exist_ok=True)
    os.makedirs(DIR_REPORTES, exist_ok=True)

    # Cargar Excel
    log(f"Cargando {ARCHIVO_EXCEL}...")
    wb = openpyxl.load_workbook(ARCHIVO_EXCEL)

    # ============ CE ============
    log("\n--- PROCESANDO CE ---")
    plan_CE = leer_plan_estudios(wb, "CE", fila_inicio=4)
    horarios_CE = leer_horarios(wb, "Grupos CE", fila_inicio=2, es_CE=True)

    sin_horario_CE = set(plan_CE.keys()) - set(horarios_CE.keys())
    if sin_horario_CE:
        log(f"Cursos sin horario en CE: {sorted(sin_horario_CE)}", "WARN")

    for cod in horarios_CE:
        if cod not in plan_CE:
            log(f"Curso en horarios pero no en plan (CE): {cod}", "ERROR")

    escribir_csv(
        os.path.join(DIR_DATA, "catalogo_CE.csv"),
        plan_CE,
        horarios_CE,
        cursos_sin_horario=sin_horario_CE,
    )

    # ============ ATI ============
    log("\n--- PROCESANDO ATI ---")
    plan_ATI = leer_plan_estudios(wb, "ATI", fila_inicio=3)
    horarios_ATI = leer_horarios(wb, "Grupos ATI", fila_inicio=3, es_CE=False)

    sin_horario_ATI = set(plan_ATI.keys()) - set(horarios_ATI.keys())
    if sin_horario_ATI:
        log(f"Cursos sin horario en ATI: {sorted(sin_horario_ATI)}", "WARN")

    for cod in horarios_ATI:
        if cod not in plan_ATI:
            log(f"Curso en horarios pero no en plan (ATI): {cod}", "ERROR")

    escribir_csv(
        os.path.join(DIR_DATA, "catalogo_ATI.csv"),
        plan_ATI,
        horarios_ATI,
        cursos_sin_horario=sin_horario_ATI,
    )

    # ============ RESUMEN ============
    log("\n" + "=" * 60)
    log("RESUMEN", "INFO")
    log(f"CE  - Cursos: {len(plan_CE)}, con horario: {len(horarios_CE)}")
    log(f"ATI - Cursos: {len(plan_ATI)}, con horario: {len(horarios_ATI)}")
    log("=" * 60)

    # Guardar reporte
    ruta_reporte = os.path.join(DIR_REPORTES, "reporte_limpieza.txt")
    with open(ruta_reporte, "w", encoding="utf-8") as f:
        f.write("\n".join(reporte_lineas))
    log(f"Reporte guardado: {ruta_reporte}", "OK")

if __name__ == "__main__":
    main()